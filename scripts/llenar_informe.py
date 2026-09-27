"""Diligencia una copia del formato U4 con los resultados exportados por el notebook (reports/tablas)."""
import copy
import json
from datetime import datetime
from pathlib import Path

import pandas as pd
from docx import Document
from docx.shared import Cm

RAIZ = Path(__file__).resolve().parents[1]
PLANTILLA = RAIZ / "U4_A1_ Formato de Entrega Trabajo Integrador.docx"
SALIDA = RAIZ / "entrega" / "U4_A1_Formato_Entrega_Completado.docx"
TABLAS = RAIZ / "reports" / "tablas"
PANTALLAZO = RAIZ / "reports" / "figures" / "app_streamlit_local.png"
META = json.loads((RAIZ / "models" / "metadata.json").read_text(encoding="utf-8"))
COMPLETAR = "[COMPLETAR]"

NEGOCIO = [
    "Proyecto académico de maestría con fines exclusivamente educativos (caso de estudio). El ICFES aplica la prueba "
    "Saber 11 a quienes terminan la educación media; el puntaje global (0–500) define el acceso a la educación superior, "
    "a becas y a créditos. Escenario planteado: una secretaría de educación de Antioquia que quisiera focalizar "
    "programas de nivelación y orientación.",
    "Con información socioeconómica y del colegio disponible antes del examen, estimar la probabilidad de que un estudiante "
    "de Antioquia (2022-2) obtenga un puntaje global ≥ 300 (alto desempeño, ≈17 % de los evaluados) y traducirla en una "
    "banda de acompañamiento (prioridad de nivelación, refuerzo focalizado o alto potencial).",
    "1) Clasificador con F1 de la clase positiva superior a la línea base (regresión logística). 2) Comparar 5 modelos "
    "clásicos y 3 ensambles con CV sobre el 70 %. 3) Identificar los factores más asociados y auditar el desempeño por "
    "grupos. 4) Desplegar el pipeline en Streamlit. Criterio de negocio: Recall ≥ 0,55, lift ≥ 2 frente al azar y F1 mayor "
    "que el de la regresión logística base.",
    "Aprendizaje supervisado de clasificación binaria con pipeline (imputación, codificación, balanceo y modelo), CV "
    "estratificada 5-fold, búsqueda aleatoria de hiperparámetros, ajuste del umbral de decisión t y evaluación en test (30 %). "
    "Capa prescriptiva: la probabilidad p define la banda (p < t/2: prioridad de nivelación; t/2 ≤ p < t: refuerzo "
    "focalizado; p ≥ t: alto potencial).",
]


def tabla(nombre):
    return pd.read_csv(TABLAS / f"{nombre}.csv")


def escribir(celda, texto):
    """Reemplaza el texto de la celda conservando el formato del primer run."""
    parrafo = celda.paragraphs[0]
    if parrafo.runs:
        parrafo.runs[0].text = str(texto)
        for run in parrafo.runs[1:]:
            run.text = ""
    else:
        parrafo.add_run(str(texto))


def llenar(tbl, filas):
    """Usa la fila vacía de la plantilla como molde y agrega una fila por registro."""
    molde = copy.deepcopy(tbl.rows[1]._tr)
    for i, valores in enumerate(filas):
        if i > 0:
            tbl._tbl.append(copy.deepcopy(molde))
        for celda, valor in zip(tbl.rows[i + 1].cells, valores):
            escribir(celda, valor)


def f4(x):
    return f"{x:.4f}"


def main():
    doc = Document(PLANTILLA)
    T = doc.tables
    if doc.paragraphs[0].text.startswith("Nota para diseño"):   # Nota interna de la plantilla
        doc.paragraphs[0]._element.getparent().remove(doc.paragraphs[0]._element)

    info = ["Predicción de alto desempeño en la prueba Saber 11 en Antioquia (2022-2) con CRISP-DM",
            f"{COMPLETAR} Nombre(s) del (los) integrante(s)", "Clasificación (binaria)",
            "Datos Abiertos Colombia — ICFES: Resultados únicos Saber 11 (kgxf-xxbe)", META["fuente"],
            f"{COMPLETAR} Enlace público de Google Colab (notebooks/02_Entregable_Proyecto_Integrador_Saber11.ipynb)",
            f"{COMPLETAR} URL del repositorio GitHub",
            f"{COMPLETAR} URL pública en Streamlit.io"]
    for fila, valor in zip(T[0].rows, info):
        escribir(fila.cells[1], valor)
    for fila, valor in zip(T[1].rows, NEGOCIO):
        escribir(fila.cells[1], valor)
    escribir(T[2].rows[1].cells[2], "Regresión Logística, KNN, Árbol de Decisión, SVM (RBF), Random Forest; "
                                    "ensambles: Votación suave, Bagging, XGBoost")
    escribir(T[2].rows[1].cells[3], "F1 de la clase positiva (alto desempeño)")

    llenar(T[3], tabla("diccionario_datos")[["variable", "descripcion", "tipo", "rol"]].values.tolist())
    llenar(T[4], tabla("reglas_numericas")[["variable", "rango_esperado"]].values.tolist())
    llenar(T[5], tabla("reglas_categoricas")[["variable", "categorias_validas"]].values.tolist())
    llenar(T[6], tabla("seleccion_variables")[["variable", "justificacion"]].values.tolist())
    for fila, valor in zip(T[7].rows[1:], tabla("descripcion_estadistica")["resultado"]):
        escribir(fila.cells[1], valor)
    # El perfilamiento describe los datos antes de limpiar: se aclara cuántos registros quedan tras la limpieza (3.3)
    errores = tabla("limpieza_errores")
    eliminados = errores.loc[errores["accion"].str.startswith("Eliminad"), "registros_afectados"].sum()
    originales = int(tabla("descripcion_estadistica")["resultado"][0].replace(",", ""))
    escribir(T[7].rows[1].cells[1], f"{originales:,} originales ({originales - eliminados:,} tras eliminar duplicados "
                                    "y resultados no oficiales)")
    llenar(T[8], [[r.tipo_error, f"{r.accion} ({r.registros_afectados:,} registros)"]
                  for r in tabla("limpieza_errores").itertuples()])
    llenar(T[9], [[r.variable, f"{r.pct_nulos:.2f} %", r.metodo] for r in tabla("limpieza_nulos").itertuples()])
    llenar(T[10], tabla("codificacion")[["variable", "metodo"]].values.tolist())
    llenar(T[11], tabla("reduccion_dimensiones")[["tecnica", "justificacion"]].values.tolist())
    llenar(T[12], tabla("balanceo")[["tecnica", "justificacion"]].values.tolist())
    llenar(T[13], [[r.modelo, r.hiperparametros_iniciales, r.metrica, f4(r.cv_mean), f4(r.cv_std)]
                   for r in tabla("modelos_clasicos").itertuples()])
    llenar(T[14], tabla("modelos_ensamble")[["modelo", "tipo", "desempeno", "mejora_vs_clasico"]].values.tolist())
    llenar(T[15], [[r.modelo, r.hiperparametro, r.que_controla,
                    f"{r.rango_probado}. Razón: {r.razon_rango}. Efecto: {r.efecto}", r.valor_final]
                   for r in tabla("ajuste_hiperparametros").itertuples()])
    llenar(T[16], [[r.modelo, f4(r.accuracy), f4(r.precision), f4(r.recall), f4(r.f1)]
                   for r in tabla("resultados_test").itertuples()])
    llenar(T[17], [["No aplica: el problema es de clasificación", "—", "—", "—"]])
    llenar(T[18], tabla("modelo_seleccionado")[["modelo_final", "justificacion"]].values.tolist())
    llenar(T[19], tabla("predicciones_nuevas")[["registro", "prediccion", "interpretacion"]].values.tolist())
    despliegue = [f"Sí — models/pipeline_saber11.joblib ({META['modelo_final']})",
                  "Sí — verificada en local (streamlit run app.py) y con pruebas automáticas",
                  f"{COMPLETAR} Sí/No tras desplegar en Streamlit.io",
                  "Sí (captura local) — reemplazar por la de la URL pública" if PANTALLAZO.exists() else COMPLETAR]
    for fila, valor in zip(T[20].rows, despliegue):
        escribir(fila.cells[1], valor)

    for parrafo in doc.paragraphs:
        if parrafo.text.startswith("Adjuntar reporte HTML"):
            parrafo.add_run(" Archivo adjunto: reports/perfilamiento_saber11_antioquia.html")
    if PANTALLAZO.exists():
        doc.add_paragraph("Pantallazo de la aplicación Streamlit (ejecución local):")
        doc.add_picture(str(PANTALLAZO), width=Cm(16))

    # Propiedades del archivo: la plantilla trae como autores a quienes la diseñaron; se dejan en blanco para que
    # Word registre a quien entrega al guardar
    propiedades = doc.core_properties
    propiedades.author, propiedades.last_modified_by = "", ""
    propiedades.title = "Proyecto Integrador CRISP-DM — Predicción de alto desempeño en Saber 11 (Antioquia 2022-2)"
    propiedades.created = propiedades.modified = datetime.now()
    propiedades.revision = 1

    SALIDA.parent.mkdir(parents=True, exist_ok=True)
    doc.save(SALIDA)
    print("Informe generado:", SALIDA)


if __name__ == "__main__":
    main()
